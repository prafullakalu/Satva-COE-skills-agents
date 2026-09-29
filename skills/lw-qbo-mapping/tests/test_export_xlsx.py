"""
Tests for export_xlsx.py -- the inventory-focused 2-sheet workbook (Matched Products /
Unmatched - Needs Review). Verifies the workbook actually contains the expected rows in the
right sheet, not just that save succeeded. Skips cleanly if openpyxl isn't installed.
"""
import sys, os, unittest, tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from normalize import SourceRecord, TargetRecord
from mapping_engine import match_product_identity, find_orphaned_products


def src(sku, name, price=None, qty=None):
    return SourceRecord(source_system="linnworks", entity_type="sku", source_id=sku, name=name,
                         amount=price, sku=sku, qty=qty)


def tgt(target_id, sku, name, price=None, qty=None, active=True, asset="Inventory Asset",
        income="Sales of Product Income", expense="Cost of Goods Sold"):
    return TargetRecord(target_system="qbo", entity_type="item", target_id=target_id, name=name,
                         sku=sku, unit_price=price, qty=qty, active=active,
                         metadata={"AssetAccountRef": {"name": asset}, "IncomeAccountRef": {"name": income},
                                   "ExpenseAccountRef": {"name": expense}})


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
class TestInventoryExport(unittest.TestCase):

    def test_matched_row_includes_coa_and_verification(self):
        from export_xlsx import build_inventory_workbook

        s = src("0051", "Two Cent Coin", price=33, qty=100)
        t = tgt("i1", "0051", "Two Cent Coin", price=33.5, qty=100)  # price within tolerance, qty exact
        result = match_product_identity(s, [t])
        self.assertEqual(result.status, "MAPPED")

        wb = build_inventory_workbook(sources=[s], targets=[t], results=[result])
        ws = wb["Matched Products"]
        self.assertEqual(ws.max_row, 2)  # header + 1 match
        row = [c.value for c in ws[2]]
        self.assertEqual(row[0], "0051")               # Linnworks SKU
        self.assertEqual(row[2], "0051")                # QBO SKU
        self.assertIn("Asset: Inventory Asset", row[4])  # COA column
        self.assertEqual(row[6], "MATCH")                # Price Match (within 5%)
        self.assertEqual(row[9], "MATCH")                # Qty Match (exact)

    def test_price_and_qty_mismatch_flagged(self):
        from export_xlsx import build_inventory_workbook

        s = src("0051", "Two Cent Coin", price=33, qty=100)
        t = tgt("i1", "0051", "Two Cent Coin", price=999, qty=5)
        result = match_product_identity(s, [t])
        wb = build_inventory_workbook(sources=[s], targets=[t], results=[result])
        ws = wb["Matched Products"]
        row = [c.value for c in ws[2]]
        self.assertEqual(row[6], "MISMATCH")  # price way off
        self.assertEqual(row[9], "MISMATCH")  # qty way off

    def test_missing_price_or_qty_is_not_fabricated_as_match(self):
        from export_xlsx import build_inventory_workbook

        s = src("0051", "Two Cent Coin", price=None, qty=None)
        t = tgt("i1", "0051", "Two Cent Coin", price=33, qty=100)
        result = match_product_identity(s, [t])
        wb = build_inventory_workbook(sources=[s], targets=[t], results=[result])
        row = [c.value for c in wb["Matched Products"][2]]
        self.assertEqual(row[6], "N/A")
        self.assertEqual(row[9], "N/A")

    def test_unmatched_lands_on_second_sheet_with_reason(self):
        from export_xlsx import build_inventory_workbook

        s = src("9999", "Mystery Coin", price=10)
        t = tgt("i1", "OTHER", "Unrelated Badge", price=5)
        result = match_product_identity(s, [t])  # no SKU/name overlap -> UNMAPPED
        self.assertEqual(result.status, "UNMAPPED")

        wb = build_inventory_workbook(sources=[s], targets=[t], results=[result])
        ws = wb["Matched Products"]
        self.assertEqual(ws.max_row, 1)  # header only, nothing matched
        ws2 = wb["Unmatched - Needs Review"]
        self.assertEqual(ws2.max_row, 2)
        row = [c.value for c in ws2[2]]
        self.assertEqual(row[0], "9999")
        self.assertEqual(row[4], "Missing in QBO")  # coarse status, not raw UNMAPPED

    def test_orphaned_qbo_item_appears_on_unmatched_sheet(self):
        from export_xlsx import build_inventory_workbook

        s = src("0051", "Two Cent Coin", price=33)
        matched = tgt("i1", "0051", "Two Cent Coin", price=33)
        orphan = tgt("i2", "9999-QBO", "Unrelated Badge", price=5)
        result = match_product_identity(s, [matched, orphan])
        orphans = find_orphaned_products([matched, orphan], [result])

        wb = build_inventory_workbook(sources=[s], targets=[matched, orphan],
                                       results=[result], orphan_results=orphans)
        ws2 = wb["Unmatched - Needs Review"]
        statuses = [ws2.cell(row=r, column=5).value for r in range(2, ws2.max_row + 1)]
        self.assertIn("Missing in Linnworks", statuses)  # coarse status for an ORPHANED QBO item

    def test_save_and_reload_roundtrip(self):
        from export_xlsx import save_inventory_workbook
        s = src("0051", "Two Cent Coin", price=33, qty=10)
        t = tgt("i1", "0051", "Two Cent Coin", price=33, qty=10)
        result = match_product_identity(s, [t])
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "test.xlsx")
            save_inventory_workbook(path, sources=[s], targets=[t], results=[result])
            self.assertTrue(os.path.exists(path))
            import openpyxl
            wb2 = openpyxl.load_workbook(path)
            self.assertEqual(set(wb2.sheetnames), {"Matched Products", "Unmatched - Needs Review"})


if __name__ == "__main__":
    unittest.main()
