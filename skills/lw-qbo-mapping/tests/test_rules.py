"""Rules store (Approved Mappings sheet) tests: load / reject / suggest / revoke."""
import sys, os, unittest, tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from openpyxl import load_workbook
import workbook as W
from workbook import DEC_HEADERS, create_workbook, update_workbook
from rules import is_rejected, load_rules, revoke_rule, suggest_from_approvals

COL = {h: i + 1 for i, h in enumerate(DEC_HEADERS)}


def dec(sid, name, tid="T1", tname="Amazon Sales"):
    return {"key": sid, "kind": "channel", "what_we_found": name, "suggestion": tname, "evidence": [], "value": 1,
            "source_id": sid, "source_name": name, "target_id": tid}


def pick(path, key, val):
    wb = load_workbook(path)
    ws = wb[W.DEC]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, COL["Key"]).value == key:
            ws.cell(r, COL["Your Decision"], val)
    wb.save(path)


class RulesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "coa.xlsx")
        create_workbook(self.path, "x")
        base = {"client": "x", "period": "p", "health_pct": 90.0, "wrongs": []}
        update_workbook(self.path, dict(base, decisions=[
            dec("uk", "Amazon UK"), dec("us", "Amazon US"), dec("eb", "eBay UK", "T2", "eBay Sales")]))
        pick(self.path, "uk", "Approve")
        pick(self.path, "us", "Approve")
        pick(self.path, "eb", "Reject")
        update_workbook(self.path, dict(base, decisions=[]))

    def tearDown(self):
        self.tmp.cleanup()

    def test_load_and_reject(self):
        r = load_rules(self.path)
        self.assertEqual(r["explicit"], {"uk": "T1", "us": "T1"})
        self.assertTrue(is_rejected(r, "eb", "T2"))
        self.assertFalse(is_rejected(r, "uk", "T1"))

    def test_suggest_needs_real_overlap_and_is_text_only(self):
        r = load_rules(self.path)
        s = suggest_from_approvals(r, [{"source_id": "de", "source_name": "Amazon DE"},
                                       {"source_id": "x", "source_name": "Etsy UK"}])
        self.assertEqual(len(s), 1)  # 'UK'-style short tokens never link Etsy UK to anything
        self.assertEqual((s[0]["source_id"], s[0]["target_id"]), ("de", "T1"))
        self.assertIn("'Amazon UK' and 'Amazon US' to 'Amazon Sales'; 'Amazon DE' looks the same", s[0]["text"])
        self.assertNotIn("de", r["explicit"])  # never auto-applied

    def test_revoke_is_audited(self):
        self.assertTrue(revoke_rule(self.path, "uk"))
        self.assertNotIn("uk", load_rules(self.path)["explicit"])
        wb = load_workbook(self.path)
        self.assertEqual([wb[W.APP].cell(r, 9).value for r in range(2, wb[W.APP].max_row + 1)].count("Revoked"), 1)
        self.assertTrue(any("Revoked" in str(c.value) for c in wb[W.LOG]["C"]))
        self.assertFalse(revoke_rule(self.path, "nope"))

    def test_rules_feed_engine(self):
        from mapping_engine import propose_mapping
        from normalize import SourceRecord, TargetRecord
        import inspect
        self.assertIn("explicit_rules", inspect.signature(propose_mapping).parameters)
        self.assertEqual(load_rules(self.path)["explicit"]["uk"], "T1")


if __name__ == "__main__":
    unittest.main()
