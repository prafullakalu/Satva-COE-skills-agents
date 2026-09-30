"""End-to-end driver test. SYNTHETIC data only -- mirrors the Coins of America lessons at small scale."""
import json, os, sys, tempfile, unittest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import run_analysis as ra, workbook
from openpyxl import load_workbook

LW = [{"SKU": "A100", "Name": "Silver Eagle", "Qty": 5}, {"SKU": "B200", "Name": "Morgan Dollar", "Qty": 2},
      {"SKU": "C300", "Name": "Peace Dollar", "Qty": 1}, {"SKU": "D400", "Name": "Buffalo Nickel Roll", "Qty": 9}]
# QBO "SKU" column is really the item Id; the Name holds the Linnworks SKU (Coins of America lesson).
QBO = [{"SKU": "1", "Name": "A100", "Qty": 5}, {"SKU": "2", "Name": "Morgan Dollar", "Qty": 2},
       {"SKU": "3", "Name": "Ghost Item", "Qty": 1}]


def dump(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f)


class TestRun(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        os.environ["LW_QBO_HOME"] = self.tmp
        self.lw, self.qbo = (os.path.join(self.tmp, n) for n in ("lw.json", "qbo.json"))
        dump(self.lw, LW); dump(self.qbo, QBO)
        self.wb = os.path.join(self.tmp, "Acme", "workbook.xlsx")

    def run_cli(self):
        return ra.main(["--client", "Acme", "--lw", self.lw, "--qbo", self.qbo])

    def test_refuses_one_sided_run(self):
        self.assertEqual(ra.main(["--client", "Acme", "--lw", self.lw]), 2)
        self.assertFalse(os.path.exists(self.wb))
        dump(self.qbo, [])
        self.assertEqual(self.run_cli(), 2)

    def test_first_run_builds_workbook_outside_repo(self):
        self.assertEqual(self.run_cli(), 0)
        self.assertTrue(os.path.exists(self.wb))
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "Acme", "summary.csv")))
        sheets = load_workbook(self.wb).sheetnames
        for s in ("Start Here", "Needs Your Decision", "Only in QBO", "Only in Linnworks"):
            self.assertIn(s, sheets)

    def test_approval_is_learned_next_run(self):
        self.run_cli()
        wbk = load_workbook(self.wb); ws = wbk["Needs Your Decision"]
        hdr = [c.value for c in ws[1]]
        # 'Morgan Dollar' (exact name, different code) is asked; 'Peace Dollar' vs 'Morgan Dollar'
        # shares one word only, so it must NOT be suggested.
        self.assertEqual(ws.max_row, 2)
        ws.cell(2, hdr.index("Your Decision") + 1, "Approve")
        wbk.save(self.wb)
        self.run_cli()
        run, _ = ra.analyse(ra.read_rows(self.lw), ra.read_rows(self.qbo), self.wb, "Acme", "p")
        self.assertEqual(run["decisions"], [], "approved item must not be asked again")
        self.assertEqual(len(run["matched"]), 2)  # A100 by name-as-code + approved Morgan Dollar
        self.assertIn("Peace Dollar", [r["Linnworks name"] for r in run["only_in_lw"]])
    def _decide(self, key_col_value, **cols):
        wbk = load_workbook(self.wb); ws = wbk["Needs Your Decision"]
        hdr = [c.value for c in ws[1]]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, hdr.index("Key") + 1).value == key_col_value:
                for h, v in cols.items():
                    ws.cell(r, hdr.index(h.replace("_", " ")) + 1, v)
        wbk.save(self.wb)

    def analyse(self):
        return ra.analyse(ra.read_rows(self.lw), ra.read_rows(self.qbo), self.wb, "Acme", "p")[0]

    def test_rejected_match_is_not_counted_matched_or_healthy(self):
        dump(self.lw, [{"SKU": "ab-100", "Name": "Alpha", "Qty": 1}])
        dump(self.qbo, [{"SKU": "1", "Name": "AB100", "Qty": 1}])
        self.run_cli()
        self._decide("ab-100", Your_Decision="Reject")
        run = self.analyse()
        self.assertEqual(run["matched"], [])
        self.assertEqual(run["health_pct"], 0)
        self.assertEqual(run["decisions"], [])

    def test_change_to_by_item_name_is_applied(self):
        self.run_cli()
        self._decide("B200", Your_Decision="Change to...", Change_to="Ghost Item")
        run = self.analyse()
        self.assertIn("Morgan Dollar", [m["Linnworks name"] for m in run["matched"]])
        self.assertNotIn("Ghost Item", [o["QuickBooks name"] for o in run["only_in_qbo"]])

    def test_duplicates_and_ambiguous_are_visible_not_lost(self):
        dump(self.lw, [{"SKU": "D1", "Name": "Dup", "Qty": 1}, {"SKU": "D1", "Name": "Dup again", "Qty": 1},
                       {"SKU": "E1", "Name": "Eagle", "Qty": 1}])
        dump(self.qbo, [{"SKU": "1", "Name": "E1", "Qty": 1}, {"SKU": "2", "Name": "E1", "Qty": 1}])
        run = self.analyse()
        self.assertTrue(any(w["key"] == "duplicate_lw_codes" for w in run["wrongs"]))
        self.assertIn("Several possible matches in QuickBooks", [r["Why"] for r in run["only_in_lw"]])
        self.assertEqual(run["only_in_qbo"], [], "tied candidates are not 'only in QuickBooks'")

    def test_two_clients_never_share_a_folder(self):
        ra.client_dir("Coins of America")
        with self.assertRaises(ra.PreflightError):
            ra.client_dir("Coins-of-America!")
        self.assertNotEqual(ra.client_dir("日本"), ra.client_dir("Müller"))


if __name__ == "__main__":
    unittest.main()
