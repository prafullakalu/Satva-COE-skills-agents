"""
Inventory-focused xlsx export: Linnworks SKU/product <-> QBO SKU/product/COA, for inventory
management specifically. Two sheets only, per requirement:

  1. "Matched Products" -- product identity confirmed (same SKU/name on both sides), showing the
     QBO chart-of-accounts object the item posts to, PLUS a post-match verification: does the
     price agree, does the on-hand quantity agree. Price/qty are checked ONLY after identity is
     confirmed -- they never decide the match itself (see mapping_engine.match_product_identity).
  2. "Unmatched - Needs Review" -- everything else.

Both sheets show a single coarse "Status" column with exactly four values (COARSE_STATUS below),
not the engine's finer-grained taxonomy and not a free-text reason:
  - Matched              -- confirmed on sheet 1
  - Missing in QBO        -- the Linnworks product has no QBO item at all (was UNMAPPED)
  - Missing in Linnworks  -- the QBO item has no Linnworks product at all (was ORPHANED)
  - Unmatched             -- some candidate was found (by name, or with a data-quality issue like
                             a duplicate SKU or an inactive target) but couldn't be confirmed
                             (was AMBIGUOUS/UNVERIFIED/INACTIVE/INVALID)

Pure formatting over already-computed data -- no matching decisions happen here. The match
decision (mapping_engine.match_product_identity) and the price/qty comparison
(_compare_value below) are the only two decision points; both are exposed so nothing is a black box.
"""
from __future__ import annotations
from normalize import SourceRecord, TargetRecord
from mapping_engine import MappingResult

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as e:
    raise ImportError("export_xlsx.py requires openpyxl (pip install openpyxl).") from e

HEADER_FILL = "305496"
HEADER_FONT = Font(color="FFFFFF", bold=True)
GREEN, YELLOW, RED, GREY = "C6EFCE", "FFEB9C", "FFC7CE", "D9D9D9"

# Coarse 4-value status vocabulary shown to the reader, in place of the engine's finer-grained
# taxonomy (reference/gap-taxonomy.md) and any free-text reason. Both sheets use only these four.
STATUS_FILL = {"Matched": GREEN, "Unmatched": YELLOW, "Missing in QBO": RED, "Missing in Linnworks": RED}
VERIFY_FILL = {"MATCH": GREEN, "MISMATCH": RED, "N/A": GREY}

# raw mapping_engine status -> coarse status. UNMAPPED means the Linnworks product has no QBO
# counterpart at all (missing FROM QBO); ORPHANED means the QBO item has no Linnworks counterpart
# (missing FROM Linnworks). Everything else found SOME candidate but couldn't confirm it.
COARSE_STATUS = {
    "MAPPED": "Matched",
    "UNMAPPED": "Missing in QBO",
    "ORPHANED": "Missing in Linnworks",
    "AMBIGUOUS": "Unmatched",
    "UNVERIFIED": "Unmatched",
    "INACTIVE": "Unmatched",
    "INVALID": "Unmatched",
}

QTY_TOLERANCE_ABS = 0     # units difference allowed before flagging a mismatch (default: exact)
PRICE_TOLERANCE_PCT = 5.0  # % difference allowed before flagging a mismatch


def _write_header(ws, headers: list[str]):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _autosize(ws, min_width=10, max_width=60):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=min_width)
        ws.column_dimensions[col_letter].width = max(min_width, min(max_width, length + 2))


def _coa_from_target(t: TargetRecord) -> str:
    """Extract the QBO chart-of-accounts object(s) this item posts to, from the raw Item metadata.
    An inventory item typically carries three: Income (sales), Expense/COGS, and Asset (inventory).
    """
    md = t.metadata or {}
    parts = []
    asset = (md.get("AssetAccountRef") or {}).get("name")
    income = (md.get("IncomeAccountRef") or {}).get("name")
    expense = (md.get("ExpenseAccountRef") or {}).get("name")
    if asset:
        parts.append(f"Asset: {asset}")
    if income:
        parts.append(f"Income: {income}")
    if expense:
        parts.append(f"COGS/Expense: {expense}")
    return " | ".join(parts) if parts else (t.subtype or "-")


def _compare_value(a, b, tolerance_pct=None, tolerance_abs=None) -> str:
    """Returns MATCH / MISMATCH / N/A. Never fabricates a verdict when either side is missing."""
    if a is None or b is None:
        return "N/A"
    if tolerance_abs is not None:
        return "MATCH" if abs(a - b) <= tolerance_abs else "MISMATCH"
    base = max(abs(a), abs(b))
    if base == 0:
        return "MATCH" if a == b else "MISMATCH"
    diff_pct = abs(a - b) / base * 100
    return "MATCH" if diff_pct <= (tolerance_pct or 0) else "MISMATCH"


def build_inventory_workbook(
    sources: list[SourceRecord],
    targets: list[TargetRecord],
    results: list[MappingResult],
    orphan_results: list[MappingResult] | None = None,
    price_tolerance_pct: float = PRICE_TOLERANCE_PCT,
    qty_tolerance_abs: float = QTY_TOLERANCE_ABS,
) -> Workbook:
    orphan_results = orphan_results or []
    source_by_id = {s.source_id: s for s in sources}
    target_by_id = {t.target_id: t for t in targets}

    wb = Workbook()

    # --- Sheet 1: Matched Products ---
    ws = wb.active
    ws.title = "Matched Products"
    _write_header(ws, [
        "Linnworks SKU", "Linnworks Product Name",
        "QBO SKU", "QBO Product Name", "QBO COA (Inventory/Income/COGS)",
        "Status",
        "Price Match", "Linnworks Price", "QBO Price",
        "Qty Match", "Linnworks Qty (owned)", "QBO Qty On Hand",
    ])
    for r in results:
        if r.status != "MAPPED":
            continue
        src = source_by_id.get(r.source_id)
        tgt = target_by_id.get(r.target_id)
        price_verdict = _compare_value(src.amount if src else None, tgt.unit_price if tgt else None,
                                        tolerance_pct=price_tolerance_pct)
        qty_verdict = _compare_value(src.qty if src else None, tgt.qty if tgt else None,
                                      tolerance_abs=qty_tolerance_abs)
        row_idx = ws.max_row + 1
        status = COARSE_STATUS[r.status]
        ws.append([
            r.source_id, r.source_name,
            (tgt.sku if tgt else "-"), r.target_name, _coa_from_target(tgt) if tgt else "-",
            status,
            price_verdict, (src.amount if src else None), (tgt.unit_price if tgt else None),
            qty_verdict, (src.qty if src else None), (tgt.qty if tgt else None),
        ])
        ws.cell(row=row_idx, column=6).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        ws.cell(row=row_idx, column=7).fill = PatternFill("solid", fgColor=VERIFY_FILL[price_verdict])
        ws.cell(row=row_idx, column=10).fill = PatternFill("solid", fgColor=VERIFY_FILL[qty_verdict])
    _autosize(ws)

    # --- Sheet 2: Unmatched / Needs Review ---
    ws2 = wb.create_sheet("Unmatched - Needs Review")
    _write_header(ws2, [
        "Linnworks SKU", "Linnworks Product Name",
        "QBO SKU", "QBO Product Name",
        "Status",
    ])
    for r in list(results) + list(orphan_results):
        if r.status == "MAPPED":
            continue
        src = source_by_id.get(r.source_id)
        tgt = target_by_id.get(r.target_id) if r.target_id else None
        row_idx = ws2.max_row + 1
        status = COARSE_STATUS[r.status]
        ws2.append([
            r.source_id or "-", (src.name if src else r.source_name) or "-",
            (tgt.sku if tgt else "-"), (tgt.name if tgt else (r.target_name or "-")),
            status,
        ])
        ws2.cell(row=row_idx, column=5).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
    _autosize(ws2, max_width=60)

    return wb


def save_inventory_workbook(path: str, **kwargs) -> str:
    wb = build_inventory_workbook(**kwargs)
    wb.save(path)
    return path
