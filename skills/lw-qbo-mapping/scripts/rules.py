"""
The 'learning' store, backed by the workbook's Approved Mappings sheet (no JSON for the accountant).
load_rules(path)["explicit"] is usable directly as explicit_rules for mapping_engine.propose_mapping.
"""
from __future__ import annotations
from openpyxl import load_workbook
from mapping_engine import _tokens
from workbook import APP, APP_HEADERS, _log, _rows, _save, _sheet, backup


def load_rules(path):
    """{'explicit': {source_id: target_id}, 'rejected': {(source_id, target_id)}, 'names': {source_id: (src_name, tgt_name)}}"""
    ws = _sheet(load_workbook(path, data_only=True), APP, APP_HEADERS)
    rules = {"explicit": {}, "rejected": set(), "names": {}}
    for _, v in _rows(ws, APP_HEADERS):
        if v["Status"] != "Active":
            continue
        sid, tid = str(v["Source ID"]), str(v["QBO Target ID"])
        if v["Decision"] == "Rejected":
            rules["rejected"].add((sid, tid))
        else:
            rules["explicit"][sid] = tid
            rules["names"][sid] = (v["Source Name"], v["QBO Target Name"])
    return rules


def is_rejected(rules, source_id, target_id):
    return (str(source_id), str(target_id)) in rules["rejected"]


def _get(o, k):
    return o.get(k) if isinstance(o, dict) else getattr(o, k, None)


def _shared(a, b):
    # ponytail: tokens <=2 chars ignored so country codes (uk/us/de) never link 'Amazon UK' with 'eBay UK'; no fuzzy matching.
    return {t for t in _tokens(a) & _tokens(b) if len(t) > 2}


def suggest_from_approvals(rules, unmapped_sources):
    """Suggestions only, never auto-applied. unmapped_sources: dicts or objects with source_id, source_name."""
    out = []
    for s in unmapped_sources:
        sid, name = str(_get(s, "source_id")), _get(s, "source_name") or ""
        by_target = {}
        for asid, (aname, tname) in rules["names"].items():
            tid = rules["explicit"][asid]
            if _shared(name, aname or "") and not is_rejected(rules, sid, tid):
                by_target.setdefault(tid, (tname, []))[1].append(aname)
        if not by_target:
            continue
        tid, (tname, srcs) = max(by_target.items(), key=lambda kv: len(kv[1][1]))
        named = " and ".join(f"'{x}'" for x in srcs[:2])
        out.append({"source_id": sid, "source_name": name, "target_id": tid, "target_name": tname,
                    "text": f"You mapped {named} to '{tname}'; '{name}' looks the same. Please confirm."})
    return out


def revoke_rule(path, source_id):
    """Mark the accountant's approval for source_id Revoked (kept for audit, logged). True if one was found."""
    path = str(path)
    backup(path)
    wb = load_workbook(path)
    ws = _sheet(wb, APP, APP_HEADERS)
    hit = False
    for r, v in _rows(ws, APP_HEADERS):
        if str(v["Source ID"]) == str(source_id) and v["Decision"] == "Approved" and v["Status"] == "Active":
            ws.cell(r, 9, "Revoked")
            _log(wb, "Accountant", f"Revoked mapping '{v['Source Name']}' -> '{v['QBO Target Name']}'", source_id)
            hit = True
    if hit:
        _save(wb, path)
    return hit
