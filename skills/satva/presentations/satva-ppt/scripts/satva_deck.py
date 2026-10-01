"""Satva slide components — cards, badges, step flows, chips, callouts.

Built on satva_pptx (clone/text helpers). These draw real PowerPoint shapes, so
everything stays editable in PowerPoint: move a card, retype a line, recolour it.

Grid (10 x 5.625in canvas, content area x 0.35 -> 9.65, y 0.90 -> 4.95):
    COL3 / COL2 / COL4  give the x positions and width for an n-across row.
    Keep anything below y=4.75 out of x>7.9 — the brand triangle lives there.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from satva_pptx import shape_by_name

BLUE = RGBColor(0x11, 0x94, 0xD2)
INK = RGBColor(0x21, 0x21, 0x21)
GREY = RGBColor(0x43, 0x43, 0x43)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xDC, 0xEB)
TINT = RGBColor(0xF2, 0xF9, 0xFD)
FONT = "Mulish"   # NOT "Muli" — the old name is installed nowhere

# n-across column layouts: (x_positions, width)
BAND = (1.85, 4.40)   # vertical content band: below the kicker, above the footer
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)

COL2 = ([0.35, 5.10], 4.55)
COL3 = ([0.35, 3.51, 6.67], 2.98)
COL4 = ([0.35, 2.70, 5.05, 7.40], 2.25)
COL5 = ([0.35, 2.22, 4.09, 5.96, 7.83], 1.78)


_GROUP = [0]


def _tag(shape, new_group=False):
    """Name shapes A<group>_<n> so animate.ps1 can reveal one group per click."""
    if new_group:
        _GROUP[0] += 1
    _GROUP[0] = max(_GROUP[0], 1)
    shape.name = "A%d_%d" % (_GROUP[0], id(shape) % 100000)
    return shape


def new_group():
    _GROUP[0] += 1


# --------------------------------------------------------------------- canvas
def blank(prs, template, title, kicker=None):
    """Clone `template` and strip it back to background + title + kicker,
    leaving the whole content area free for components."""
    from satva_pptx import clone_slide, set_text

    s = clone_slide(prs, template)
    keep = {"Google Shape;64;p14", "Google Shape;65;p14", "Google Shape;66;p14"}
    for shp in list(s.shapes):
        if shp.name not in keep:
            shp._element.getparent().remove(shp._element)
    set_text(shape_by_name(s, "Google Shape;65;p14"), title)
    if kicker is not None:
        set_text(shape_by_name(s, "Google Shape;66;p14"), kicker)
    return s


# ----------------------------------------------------------------------- text
def _tb(slide, x, y, w, h):
    box = _tag(slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def _write(tf, lines, size, color, bold=False, space=4, align=PP_ALIGN.LEFT,
           bullet=False):
    r"""`lines` may be strings or (text, bold) pairs. `**lead**` bolds the lead.

    A "\n" inside a run is invalid DrawingML and makes PowerPoint declare the
    whole deck corrupt, so multi-line strings become separate paragraphs here."""
    first = True
    lines = [part for line in lines
             for part in (line.split("\n") if isinstance(line, str) else [line])]
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = 1.12
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "142875")
            pPr.set("indent", "-142875")
        chunks = []
        if isinstance(line, str) and line.startswith("**") and "**" in line[2:]:
            lead, rest = line[2:].split("**", 1)
            chunks = [(lead, True), (rest, False)]
        else:
            chunks = [(line, bold)]
        for text, b in chunks:
            if text == "":
                continue
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.name = FONT
            r.font.color.rgb = color


def text(slide, x, y, w, h, lines, size=12, color=GREY, bold=False, space=4,
         align=PP_ALIGN.LEFT, bullet=False):
    box, tf = _tb(slide, x, y, w, h)
    _write(tf, lines, size, color, bold, space, align, bullet)
    return box


def footer(slide, line, y=None):
    """The one-line takeaway across the bottom, brand blue, centred.

    With no `y` it sits 0.32in under whatever component row was drawn last, so
    it hugs the content instead of stranding a band of white above it."""
    if y is None:
        y = min(4.86, getattr(slide, "_bottom", 4.50) + 0.32)
    new_group()
    return text(slide, 0.35, y, 9.22, 0.4, [line], size=12, color=BLUE,
                align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------- cards
def card(slide, x, y, w, h, title, lines, badge=None, tint=False, size=11,
         badge_fill=None):
    """A rounded content card: optional badge, blue title, grey body lines."""
    box = _tag(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(x), Inches(y), Inches(w), Inches(h)), True)
    box.adjustments[0] = 0.06
    box.fill.solid()
    box.fill.fore_color.rgb = TINT if tint else WHITE
    box.line.color.rgb = LINE
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    box.text_frame.text = ""

    tx = x + 0.20
    if badge is not None:
        dot = _tag(slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.20),
                                          Inches(y + 0.18), Inches(0.36), Inches(0.36)))
        dot.fill.solid()
        dot.fill.fore_color.rgb = badge_fill or BLUE
        dot.line.fill.background()
        dot.shadow.inherit = False
        tf = dot.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _write(tf, [str(badge)], 12, WHITE, True, 0, PP_ALIGN.CENTER)
        tx = x + 0.66

    text(slide, tx, y + 0.22, w - (tx - x) - 0.18, 0.32, [title], size=12.5,
         color=BLUE, bold=True)
    if lines:
        text(slide, x + 0.20, y + 0.66, w - 0.40, h - 0.84, lines, size=size,
             color=GREY, space=5)
    return box


def cards(slide, defs, y=None, h=2.10, cols=COL3, tint=False, size=11,
          badge_fills=None):
    """A row of cards. `defs` is [(title, [lines]), ...] or [(title, lines, badge)].

    `y=None` centres the row in the content band, so short cards sit optically
    between the kicker and the footer instead of leaving dead space at the bottom.
    Size the card to its text — 0.66in for the title row, ~0.20in per WRAPPED
    body line, 0.18in padding — then let this centre it."""
    xs, w = cols
    if y is None:
        y = BAND[0] + max(0.0, (BAND[1] - BAND[0] - h) / 2)
    out = []
    for i, d in enumerate(defs):
        title, lines = d[0], d[1]
        badge = d[2] if len(d) > 2 else None
        fill = badge_fills[i] if badge_fills else None
        out.append(card(slide, xs[i], y, w, h, title, lines, badge, tint, size,
                        badge_fill=fill))
    slide._bottom = y + h
    return out


# ----------------------------------------------------------------------- flow
def steps(slide, items, x=0.35, y=1.95, w=9.22, row=0.62, size=11.5):
    """Vertical numbered flow: blue circle + bold label + grey detail."""
    for i, (label, detail) in enumerate(items):
        top = y + i * row
        dot = _tag(slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(top),
                                          Inches(0.38), Inches(0.38)), True)
        dot.fill.solid()
        dot.fill.fore_color.rgb = BLUE
        dot.line.fill.background()
        dot.shadow.inherit = False
        tf = dot.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _write(tf, [str(i + 1)], 12, WHITE, True, 0, PP_ALIGN.CENTER)
        text(slide, x + 0.55, top + 0.02, w - 0.55, 0.30, [label], size=size + 0.5,
             color=BLUE, bold=True)
        if detail:
            text(slide, x + 0.55, top + 0.29, w - 0.55, 0.28, [detail], size=size,
                 color=GREY)
    slide._bottom = y + len(items) * row


def chevrons(slide, labels, y=2.05, h=0.72, x=0.35, w=9.22, gap=0.06):
    """Horizontal chevron pipeline — for a short, ordered process."""
    n = len(labels)
    cw = (w - gap * (n - 1)) / n
    for i, label in enumerate(labels):
        shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sh = _tag(slide.shapes.add_shape(shape, Inches(x + i * (cw + gap)), Inches(y),
                                         Inches(cw), Inches(h)), True)
        sh.fill.solid()
        sh.fill.fore_color.rgb = BLUE if i % 2 == 0 else RGBColor(0x3F, 0xAD, 0xDE)
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _write(tf, [label], 11.5, WHITE, True, 0, PP_ALIGN.CENTER)
    slide._bottom = y + h


def chips(slide, items, y=2.00, cols=COL5, rows_gap=0.78, h=0.66, size=10):
    """Grid of small labelled chips — inventories, catalogs, tool lists.
    `items` is [(name, sub), ...]; wraps onto new rows automatically."""
    xs, w = cols
    per = len(xs)
    for i, (name, sub) in enumerate(items):
        cx, cy = xs[i % per], y + (i // per) * rows_gap
        sh = _tag(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx),
                                         Inches(cy), Inches(w), Inches(h)), True)
        sh.adjustments[0] = 0.14
        sh.fill.solid()
        sh.fill.fore_color.rgb = TINT
        sh.line.color.rgb = LINE
        sh.line.width = Pt(0.75)
        sh.shadow.inherit = False
        sh.text_frame.text = ""
        text(slide, cx + 0.12, cy + 0.11, w - 0.24, 0.22, [name], size=size + 1,
             color=BLUE, bold=True)
        text(slide, cx + 0.12, cy + 0.35, w - 0.24, 0.22, [sub], size=size - 1.5,
             color=GREY)
        slide._bottom = cy + h


def callout(slide, x, y, w, h, lines, size=11.5):
    """Tinted panel for a contrast block — 'the pain' vs 'the fix'."""
    sh = _tag(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h)), True)
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = TINT
    sh.line.color.rgb = LINE
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    text(slide, x + 0.22, y + 0.20, w - 0.44, h - 0.36, lines, size=size, space=6)
    slide._bottom = y + h
    return sh


def code(slide, x, y, w, h, header, lines, size=10):
    """Terminal-style block: blue header, Consolas body, tinted panel."""
    sh = _tag(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h)), True)
    sh.adjustments[0] = 0.05
    sh.fill.solid()
    sh.fill.fore_color.rgb = TINT
    sh.line.color.rgb = LINE
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    text(slide, x + 0.20, y + 0.16, w - 0.40, 0.26, [header], size=11.5,
         color=BLUE, bold=True)
    box, tf = _tb(slide, x + 0.20, y + 0.48, w - 0.40, h - 0.62)
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = "Consolas"
        r.font.color.rgb = GREY
    slide._bottom = y + h
    return sh


# ------------------------------------------------------- diagram primitives
# Cards/steps/chevrons/chips alone force every slide into the same shape. These
# four draw the parts a diagram needs — a node, an arrow, a connector line and a
# free label — so a slide can be a tree, a decision branch or an architecture.
def node(slide, x, y, w, h, label, sub=None, fill=None, ink=None, size=11.5,
         sub_size=9.5, radius=0.10, border=True, group=True):
    """A labelled box. `fill=BLUE` gives a solid brand node with white text;
    `fill=None` gives a white node with a blue label and a hairline border."""
    solid = fill is not None
    sh = _tag(slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h)), group)
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill if solid else WHITE
    if border and not solid:
        sh.line.color.rgb = LINE
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    colour = ink or (WHITE if solid else BLUE)
    _write(tf, [label], size, colour, True, 0, PP_ALIGN.CENTER)
    if sub:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = sub
        r.font.size = Pt(sub_size)
        r.font.bold = False
        r.font.name = FONT
        r.font.color.rgb = WHITE if solid else GREY
    slide._bottom = y + h
    return sh


def arrow(slide, x, y, w, h, direction="right", color=None, group=False):
    """A solid brand arrow between two nodes. direction: right|down|left|up."""
    shapes = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW,
              "left": MSO_SHAPE.LEFT_ARROW, "up": MSO_SHAPE.UP_ARROW}
    sh = _tag(slide.shapes.add_shape(shapes[direction], Inches(x), Inches(y),
                                     Inches(w), Inches(h)), group)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color or BLUE
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def rule(slide, x, y, w, h=0.015, color=None, group=False):
    """A hairline connector — the stems and cross-bars of a tree diagram.
    Pass a tall, narrow w/h for a vertical line."""
    sh = _tag(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h)), group)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color or BLUE
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def label(slide, x, y, w, lines, size=10.5, color=None, bold=False,
          align=PP_ALIGN.CENTER, space=2):
    """Free-standing caption — the text under a node or beside an arrow."""
    return text(slide, x, y, w, 0.3, lines, size=size, color=color or GREY,
                bold=bold, space=space, align=align)
