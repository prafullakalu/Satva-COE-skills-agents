"""Satva PPTX helpers — clone slides, replace text without losing formatting.

Everything here exists because python-pptx cannot duplicate a slide and
`text_frame.text = "..."` destroys the run formatting. Both are solved by
copying the existing XML and only swapping the strings inside it.

Usage:
    import sys; sys.path.insert(0, r"C:\\Users\\Admin\\.claude\\skills\\satva-ppt\\scripts")
    from satva_pptx import *
"""
import copy
import re

from pptx import Presentation  # noqa: F401  (re-exported for callers)
from pptx.oxml.ns import qn

__all__ = [
    "Presentation", "clone_slide", "shape_by_name", "set_text", "set_bullets",
    "set_steps", "set_table", "reorder_slides", "dump",
    "set_font_family", "set_transition",
]

_BOLD = re.compile(r"^\*\*(.+?)\*\*(.*)$", re.S)


# --------------------------------------------------------------------- slides
def clone_slide(prs, src, after=None):
    """Deep-copy `src` (a Slide) into a new slide at the end. Returns the new slide.

    Copies every shape and every relationship except the layout and the speaker
    notes, so background images and picture fills survive. Reorder afterwards
    with reorder_slides().

    The notes part is deliberately dropped: a notesSlide belongs to exactly one
    slide, so pointing several clones at the source's one makes PowerPoint
    refuse to open the file ("needs repair") with no other symptom.
    """
    dst = prs.slides.add_slide(src.slide_layout)
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)

    # python-pptx picks its own rIds, so copy the rels first and remap the
    # r:embed / r:id references inside the copied shapes to the new ones.
    remap = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(("slideLayout", "notesSlide")):
            continue
        target = rel.target_ref if rel.is_external else rel.target_part
        remap[rid] = dst.part.rels._add_relationship(rel.reltype, target, rel.is_external)

    ns_r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr, val in list(node.attrib.items()):
                if attr.startswith(ns_r) and val in remap:
                    node.set(attr, remap[val])
        dst.shapes._spTree.append(el)
    return dst


def reorder_slides(prs, order):
    """`order` is a list of 0-based indices into the CURRENT slide order.
    Slides not listed are dropped from the deck."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for e in ids:
        lst.remove(e)
    for i in order:
        lst.append(ids[i])


def shape_by_name(slide, name):
    for shp in slide.shapes:
        if shp.name == name:
            return shp
    raise KeyError(f"{name!r} not in {[s.name for s in slide.shapes]}")


# ----------------------------------------------------------------------- text
def _runs(p):
    return p.findall(qn("a:r"))


def _strip_extra_runs(p, keep):
    for r in _runs(p)[keep:]:
        p.remove(r)


def _set_run(r, text):
    t = r.find(qn("a:t"))
    t.text = text
    if text != text.strip():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _set_bold(r, on):
    rPr = r.find(qn("a:rPr"))
    if rPr is None:
        rPr = r.makeelement(qn("a:rPr"), {})
        r.insert(0, rPr)
    rPr.set("b", "1" if on else "0")


def _fill(p, text):
    """Put `text` into paragraph `p`, reusing its existing run formatting.
    `**Lead:** rest` bolds the lead, whatever the template paragraph did."""
    runs = _runs(p)
    if not runs:
        return
    m = _BOLD.match(text)
    if m:
        if len(runs) < 2:                      # template had a single run — split it
            clone = copy.deepcopy(runs[0])
            runs[0].addnext(clone)
            runs = _runs(p)
        _set_run(runs[0], m.group(1))
        _set_run(runs[1], m.group(2))
        _set_bold(runs[0], True)
        _set_bold(runs[1], False)
        _strip_extra_runs(p, 2)
    elif len(runs) >= 2:
        # template mixes a bold lead with plain body — plain text keeps one run, unbolded
        _set_run(runs[1], text)
        _set_bold(runs[1], False)
        p.remove(runs[0])
        _strip_extra_runs(p, 1)
    else:
        _set_run(runs[0], text)
        _strip_extra_runs(p, 1)


def set_text(shape, text):
    """Single-paragraph shapes: titles, kickers, footers."""
    paras = shape.text_frame._txBody.findall(qn("a:p"))
    _fill(paras[0], text)
    for p in paras[1:]:
        p.getparent().remove(p)


def set_bullets(shape, items, template=0, header=None, header_template=0):
    """Replace a body text box with `items`, every paragraph styled like
    the existing paragraph at index `template`. `header` (column slides) is
    written first using the style at `header_template`."""
    body = shape.text_frame._txBody
    paras = body.findall(qn("a:p"))
    tmpl = copy.deepcopy(paras[template])
    h_tmpl = copy.deepcopy(paras[header_template])
    for p in paras:
        body.remove(p)
    if header is not None:
        p = copy.deepcopy(h_tmpl)
        _fill(p, header)
        body.append(p)
    for it in items:
        p = copy.deepcopy(tmpl)
        _fill(p, it)
        body.append(p)


def set_steps(shape, pairs, head=0, body_i=1):
    """Numbered/step slides: `pairs` is [(heading, description), ...].
    Headings copy the style of paragraph `head`, descriptions of `body_i`."""
    body = shape.text_frame._txBody
    paras = body.findall(qn("a:p"))
    t_head, t_body = copy.deepcopy(paras[head]), copy.deepcopy(paras[body_i])
    for p in paras:
        body.remove(p)
    for h, d in pairs:
        ph = copy.deepcopy(t_head)
        _fill(ph, h)
        body.append(ph)
        if d:
            pd = copy.deepcopy(t_body)
            _fill(pd, d)
            body.append(pd)


def set_table(shape, rows):
    """`rows` is a list of lists of strings, sized to the existing table.
    Cell run formatting (header fill, fonts) is preserved."""
    tbl = shape.table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            paras = cell.text_frame._txBody.findall(qn("a:p"))
            lines = val.split("\n")
            tmpl = copy.deepcopy(paras[0])
            for p in paras:
                p.getparent().remove(p)
            for line in lines:
                p = copy.deepcopy(tmpl)
                _fill(p, line)
                cell.text_frame._txBody.append(p)


# ----------------------------------------------------------------------- read
def dump(path):
    prs = Presentation(path)
    for i, s in enumerate(prs.slides, 1):
        print(f"\n===== SLIDE {i} =====")
        for sh in s.shapes:
            if sh.has_table:
                print(f"  TABLE {sh.name}")
                for row in sh.table.rows:
                    print("    | " + " | ".join(c.text.replace("\n", " / ") for c in row.cells))
            elif sh.has_text_frame:
                print(f"  TXT {sh.name!r}")
                for p in sh.text_frame.paragraphs:
                    if p.text.strip():
                        print(f"     {p.text}")
            else:
                print(f"  {sh.shape_type} {sh.name!r}")


if __name__ == "__main__":
    import sys
    dump(sys.argv[1])


# ------------------------------------------------------------------- polish
def set_font_family(prs, new="Mulish", old=("Muli",)):
    """Repoint every run using a missing family at an installed one.

    The Satva template asks for "Muli" — renamed to "Mulish" in 2020, so the old
    name resolves nowhere and PowerPoint substitutes a fallback on every run.
    That shows up as uneven weight and spacing, not as a missing-font warning."""
    n = 0
    for slide in prs.slides:
        for el in slide.shapes._spTree.iter():
            for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
                if el.tag == qn(tag) and el.get("typeface") in old:
                    el.set("typeface", new)
                    n += 1
    return n


def set_transition(prs, kind="fade", speed="med"):
    """Add a slide transition to every slide (OOXML — no COM needed)."""
    from pptx.oxml.ns import nsmap
    for slide in prs.slides:
        sld = slide._element
        for existing in sld.findall(qn("p:transition")):
            sld.remove(existing)
        tr = sld.makeelement(qn("p:transition"), {"spd": speed})
        tr.append(tr.makeelement(qn("p:" + kind), {}))
        clr = sld.find(qn("p:clrMapOvr"))
        if clr is not None:
            clr.addnext(tr)
        else:
            sld.append(tr)
