"""Workbook create/update/preserve/backup/damage/locked-file tests. Synthetic data, temp dirs, no live calls."""
import sys, os, unittest, tempfile
from unittest import mock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

from openpyxl import Workbook, load_workbook
import workbook as W
from workbook import DEC, DEC_HEADERS, WorkbookDamaged, create_workbook, load_decisions, update_workbook

COL = {h: i + 1 for i, h in enumerate(DEC_HEADERS)}


def dec(key, **kw):
    d = {"key": key, "kind": "channel", "what_we_found": f"Channel {key}", "suggestion": "Amazon Sales",
         "evidence": ["name match"], "value": 100.0, "source_id": key, "source_name": f"Src {key}", "target_id": "T1"}
    d.update(kw)
    return d


def run(decisions, **kw):
    r = {"client": "Coins of America", "period": "2026-08", "run_at": "2026-09-01 10:00",
         "summary": [("Products", 10)], "decisions": decisions, "matched": [{"sku": "A", "qbo": "A"}],
         "only_in_qbo": [], "only_in_lw": [{"sku": "Z"}],
         "wrongs": [{"key": "w1", "severity": "high", "plain_text": "Bad", "money_at_stake": 5}], "health_pct": 88.0}
    r.update(kw)
    return r


def edit(path, key, **cols):
    wb = load_workbook(path)
    ws = wb[DEC]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, COL["Key"]).value == key:
            for h, v in cols.items():
                ws.cell(r, COL[h.replace("_", " ")], v)
    wb.save(path)


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "coa.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_create_has_sheets_hidden_id_dropdown(self):
        create_workbook(self.path, "Coins of America")
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, [W.START, W.DEC, W.APP, W.MATCHED, W.ONLY_Q, W.ONLY_L, W.WRONG, W.HIST,
                                         W.LOG, W.GLOSS])
        self.assertTrue(wb[DEC].column_dimensions["A"].hidden)
        self.assertEqual(wb[DEC].freeze_panes, "A2")
        self.assertEqual(len(wb[DEC].data_validations.dataValidation), 1)

    def test_update_edit_update_preserves_and_resolves(self):
        create_workbook(self.path, "Coins of America")
        update_workbook(self.path, run([dec("a"), dec("b")]))
        edit(self.path, "a", Your_Decision="Approve", Your_Notes="ok")
        edit(self.path, "b", Your_Decision="Change to...", Change_to="T9", Your_Notes="hmm")
        update_workbook(self.path, run([dec("a", suggestion="NEW TEXT"), dec("c")]))  # b vanished, c new
        ws = load_workbook(self.path)[DEC]
        rows = {ws.cell(r, COL["Key"]).value: r for r in range(2, ws.max_row + 1)}
        self.assertEqual(set(rows), {"a", "b", "c"})
        a, b, c = rows["a"], rows["b"], rows["c"]
        self.assertEqual(ws.cell(a, COL["Your Decision"]).value, "Approve")
        self.assertEqual(ws.cell(a, COL["Your Notes"]).value, "ok")
        self.assertNotEqual(ws.cell(a, COL["Our suggestion"]).value, "NEW TEXT")  # decided row keeps what he saw
        self.assertEqual(ws.cell(b, COL["Change to"]).value, "T9")
        self.assertEqual(ws.cell(a, COL["Status"]).value, "Decided")
        self.assertEqual(ws.cell(c, COL["Status"]).value, "Open")
        # Approved Mappings got both (a: approve T1, b: change T9)
        app = load_workbook(self.path)[W.APP]
        got = {(app.cell(r, 2).value, app.cell(r, 4).value) for r in range(2, app.max_row + 1)}
        self.assertEqual(got, {("a", "T1"), ("b", "T9")})
        self.assertGreaterEqual(load_workbook(self.path)[W.HIST].max_row, 3)  # header + 2 runs

    def test_change_to_number_does_not_crash_and_resolves_via_lookup(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a")]))
        edit(self.path, "a", Your_Decision="Change to...", Change_to=8)  # Excel stores a typed number as int
        self.assertEqual(W.load_decisions(self.path)[0]["change_to"], "8")
        W.apply_decisions(self.path, lambda t: ("8", "Item Eight") if t == "8" else None)
        app = load_workbook(self.path)[W.APP]
        self.assertEqual((app.cell(2, 4).value, app.cell(2, 5).value), ("8", "Item Eight"))

    def test_change_to_not_found_stays_pending_with_status_flag(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a")]))
        edit(self.path, "a", Your_Decision="Change to...", Change_to="Nope")
        W.apply_decisions(self.path, lambda t: None)
        ws = load_workbook(self.path)[DEC]
        self.assertEqual(ws.cell(2, COL["Status"]).value, W.CHECK)
        self.assertEqual(W.open_count(self.path), 1)
        self.assertEqual(load_workbook(self.path)[W.APP].max_row, 1)  # no rule created

    def test_he_can_change_his_mind_and_old_rule_is_revoked(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a")]))
        edit(self.path, "a", Your_Decision="Approve")
        W.apply_decisions(self.path)
        edit(self.path, "a", Your_Decision="Reject")
        W.apply_decisions(self.path)
        app = load_workbook(self.path)[W.APP]
        rows = [(app.cell(r, 6).value, app.cell(r, 9).value) for r in range(2, app.max_row + 1)]
        self.assertEqual(rows, [("Rejected", "Active")])  # same target: one row updated, not duplicated

    def test_decided_row_reopens_when_qbo_offers_a_different_target(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a")]))
        edit(self.path, "a", Your_Decision="Reject")
        update_workbook(self.path, run([dec("a", target_id="T2")]))
        ws = load_workbook(self.path)[DEC]
        self.assertEqual(ws.cell(2, COL["Status"]).value, "Open")
        self.assertIn(ws.cell(2, COL["Your Decision"]).value, (None, ""))

    def test_vanished_undecided_marked_resolved_not_deleted(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a")]))
        update_workbook(self.path, run([]))
        ws = load_workbook(self.path)[DEC]
        self.assertEqual(ws.cell(2, COL["Status"]).value, "Resolved")

    def test_backup_created_and_pruned(self):
        create_workbook(self.path, "x")
        for _ in range(13):
            update_workbook(self.path, run([]))
        files = os.listdir(os.path.join(self.tmp.name, "backups"))
        self.assertEqual(len(files), 10)

    def test_damaged_header_raises_plain_english(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a")]))
        wb = load_workbook(self.path)
        wb[DEC].cell(1, COL["Your Decision"], "My choice")
        wb.save(self.path)
        with self.assertRaises(WorkbookDamaged) as cm:
            load_decisions(self.path)
        self.assertIn("backups", str(cm.exception))
        with self.assertRaises(WorkbookDamaged):
            update_workbook(self.path, run([]))

    def test_load_decisions(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a"), dec("b"), dec("c")]))
        edit(self.path, "a", Your_Decision="Approve")
        edit(self.path, "b", Your_Decision="Change to...", Change_to=" T5 ", Your_Notes="n")
        got = {d["key"]: d for d in load_decisions(self.path)}
        self.assertEqual(set(got), {"a", "b"})
        self.assertEqual(got["b"], {"key": "b", "decision": "Change", "change_to": "T5", "notes": "n"})

    def test_locked_file_writes_pending(self):
        create_workbook(self.path, "x")
        real_save = Workbook.save

        def save(self_, p):
            if str(p) == self.path:
                raise PermissionError("locked")
            return real_save(self_, p)

        with mock.patch.object(Workbook, "save", save):
            out = update_workbook(self.path, run([dec("a")]))
        self.assertTrue(out.endswith("coa_pending.xlsx"))
        self.assertEqual(load_workbook(out)[DEC].cell(2, COL["Key"]).value, "a")

    def test_start_here_plain_english(self):
        create_workbook(self.path, "x")
        update_workbook(self.path, run([dec("a"), dec("b")]))
        self.assertIn("2 things need your decision, 1 problems are urgent", load_workbook(self.path)[W.START]["A3"].value)


if __name__ == "__main__":
    unittest.main()
