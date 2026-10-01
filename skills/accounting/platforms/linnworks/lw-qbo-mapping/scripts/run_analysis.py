"""
One-call run: Linnworks + QuickBooks data in -> per-client Excel workbook + summary CSV out.

    python scripts/run_analysis.py --client "Coins of America" --lw lw.xlsx --qbo qbo.xlsx [--mcp]

Inputs are .xlsx/.csv exports, or .json dumps the agent saved from the MCP tools (add --mcp when the
JSON is raw MCP tool output rather than export rows). BOTH sides are required: no report is written
from one system alone. Client data lives in ~/LW-QBO-Mapping/<client>/ (override: LW_QBO_HOME),
never inside the repo. No credentials are read or stored here -- they belong to the MCP connections.
"""
from __future__ import annotations
import argparse, csv, json, os, re, sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import normalize, mapping_engine as me, coverage, plain, rules, summary, workbook  # noqa: E402


class PreflightError(Exception):
    """Plain-English reason the run can't start."""


def read_rows(path: str) -> list[dict]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else next((v for v in data.values() if isinstance(v, list)), [])
    if ext == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb.worksheets[0].iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it, [])]
        return [dict(zip(head, r)) for r in it if any(c is not None for c in r)]
    finally:
        wb.close()


def preflight(lw_path, qbo_path):
    """Both systems must be present and non-empty, else refuse -- never report on one side only."""
    out = {}
    for label, p in (("Linnworks", lw_path), ("QuickBooks", qbo_path)):
        if not p or not os.path.exists(p):
            raise PreflightError(f"I need the {label} list to run. Please give me the {label} data first "
                                 f"(nothing was reported, because a one-sided report would be misleading).")
        rows = read_rows(p)
        if not rows:
            raise PreflightError(f"The {label} list is empty, so I stopped instead of reporting a wrong picture.")
        out[label] = rows
    return out["Linnworks"], out["QuickBooks"]


def client_dir(client: str) -> str:
    home = os.environ.get("LW_QBO_HOME") or os.path.join(os.path.expanduser("~"), "LW-QBO-Mapping")
    d = os.path.join(home, re.sub(r"[^\w]+", "-", client).strip("-_") or "client")
    os.makedirs(d, exist_ok=True)
    prof = os.path.join(d, "profile.json")
    if os.path.exists(prof):
        with open(prof, encoding="utf-8") as f:
            owner = json.load(f).get("client", "")
        if owner.casefold() != client.casefold():  # never let two clients share learned rules
            raise PreflightError(f"The folder {d} already belongs to '{owner}'. Please use a different client name "
                                 f"so '{client}' doesn't share {owner}'s approvals.")
    else:
        with open(prof, "w", encoding="utf-8") as f:
            json.dump({"client": client, "created": str(date.today())}, f, indent=2)
    return d


MIN_NAME_SIMILARITY = 0.6  # one shared word ("dollar") must not become a suggestion an accountant might approve


def _similar(a: str, b: str) -> bool:
    x, y = me._tokens(a), me._tokens(b)
    return bool(x | y) and len(x & y) / len(x | y) >= MIN_NAME_SIMILARITY


def _match(sources, targets, rule_set, sku_key):
    by_id, idx = {t.target_id: t for t in targets}, me.build_product_index(targets, sku_key)
    results, used = [], []
    for s in sources:
        t = by_id.get(rule_set["explicit"].get(s.source_id))
        if t:
            results.append(me._product_result(s, t, "MAPPED" if t.active else "INACTIVE", "HIGH", ["approved by accountant"],
                                              "Approved in the workbook", not t.active))
            used.append(s.source_id)
            continue
        r = me.match_product_identity(s, targets, sku_key=sku_key, index=idx)
        if r.target_id and rules.is_rejected(rule_set, s.source_id, r.target_id):  # he said no: not a match
            r = me._product_result(s, None, "UNMAPPED", "UNRESOLVED", ["rejected by accountant"],
                                   "You rejected this match", True)
        results.append(r)
    return results, used


def _resolver(targets):
    """What he typed under 'Change to' -> (item id, name); only when it points to exactly one QuickBooks item."""
    def find(text):
        t = text.strip().casefold()
        hits = [x for x in targets if x.target_id.casefold() == t] or [x for x in targets if x.name.casefold() == t]
        return (hits[0].target_id, hits[0].name) if len(hits) == 1 else None
    return find


LOW_URGENCY_GAPS = {"no_chart_of_accounts", "marketplace_fees_not_in_linnworks"}  # true for every product-only run
WHY = {"UNMAPPED": "Missing in QuickBooks", "AMBIGUOUS": "Several possible matches in QuickBooks",
       "INACTIVE": "Found in QuickBooks but switched off"}


def analyse(lw_rows, qbo_rows, wb_path, client, period, mcp=False):
    sources = normalize.from_stock_items(lw_rows) if mcp else normalize.from_lw_export_rows(lw_rows)
    targets = normalize.from_qbo_items(qbo_rows) if mcp else normalize.from_qbo_export_rows(qbo_rows)
    seen, dupes, uniq = set(), set(), []
    for s in sources:  # one row per Linnworks product code; repeats are reported, not double-counted
        (dupes.add(s.source_id) if s.source_id in seen else uniq.append(s))
        seen.add(s.source_id)
    sources = uniq
    resolve = _resolver(targets)
    if os.path.exists(wb_path):
        workbook.apply_decisions(wb_path, resolve)  # his latest picks count in THIS run
        rule_set = rules.load_rules(wb_path)
    else:
        rule_set = {"explicit": {}, "rejected": set()}
    info = me.detect_sku_key(sources, targets)
    results, used = _match(sources, targets, rule_set, info["sku_key_source"])
    orphans = me.find_orphaned_products(targets, results)
    amb_names = {n for r in results if r.status == "AMBIGUOUS" for n in (r.target_name or "").split(", ")}
    orphans = [o for o in orphans if o.target_name not in amb_names]  # candidates of a tie are not "missing"
    gaps = me.data_gap_findings(sources, targets) + info["data_quality"]
    args = (client, info, sources, targets, results + orphans, gaps)
    metrics, plain_metrics = summary.summary_metrics(*args), summary.summary_metrics(*args, plain=True)
    src_by, tgt_by = {s.source_id: s for s in sources}, {t.target_id: t for t in targets}

    decisions, matched, only_lw = [], [], []
    for r in results:
        s = src_by[r.source_id]
        if r.status == "MAPPED":
            t = tgt_by.get(r.target_id)
            q = "N/A" if s.qty is None or t is None or t.qty is None else ("MATCH" if s.qty == t.qty else "MISMATCH")
            matched.append({"key": r.source_id, "Linnworks product code": s.sku, "Linnworks name": s.name,
                            "QuickBooks name": r.target_name, "Linnworks qty": s.qty,
                            "QuickBooks qty": t.qty if t else None, "Quantity check": q})
        weak = r.status == "UNVERIFIED" and not _similar(s.name, r.target_name or "")
        if r.status in WHY or weak:
            only_lw.append({"key": r.source_id, "Linnworks product code": s.sku, "Linnworks name": s.name,
                            "Why": WHY.get(r.status, WHY["UNMAPPED"]),
                            "QuickBooks candidates": r.target_name if r.status != "UNMAPPED" and not weak else None})
        if r.target_id and (r.status == "UNVERIFIED" or (r.status == "MAPPED" and r.review_required)) and not weak:
            tag = "near_sku_match" if "near SKU match (normalized)" in r.evidence else \
                "exact_name_sku_differs" if "exact_name_sku_differs" in r.evidence else r.status
            decisions.append({"key": r.source_id, "kind": "product", "what_we_found": f"{s.name} (product code {s.sku})",
                              "suggestion": r.target_name, "evidence": plain.STATUS_PLAIN.get(tag, plain.STATUS_PLAIN["UNVERIFIED"])["sentence"],
                              "value": s.amount, "source_id": r.source_id, "source_name": s.name, "target_id": r.target_id})
    only_qbo = [{"key": o.target_id, "QuickBooks name": o.target_name} for o in orphans]

    wrongs = [{"key": g["gap"], "severity": "low" if g["gap"] in LOW_URGENCY_GAPS else g["severity"].lower(),
               "plain_text": plain.gap_text(g), "money_at_stake": 0} for g in gaps]
    mism = sum(1 for m in matched if m["Quantity check"] == "MISMATCH")
    if mism:
        wrongs.append({"key": "qty_mismatch", "severity": "medium", "money_at_stake": 0,
                       "plain_text": f"{mism} matched products show a different quantity in each system."})
    if dupes:
        wrongs.append({"key": "duplicate_lw_codes", "severity": "high", "money_at_stake": 0,
                       "plain_text": f"{len(dupes)} product codes appear more than once in Linnworks; only the first of each was checked."})
    confirmed = sum(1 for r in results if r.status == "MAPPED" and not r.review_required)
    run = {"client": client, "period": period, "run_at": str(date.today()), "summary": plain_metrics,
           "decisions": decisions, "matched": matched, "only_in_qbo": only_qbo, "only_in_lw": only_lw,
           "wrongs": wrongs, "health_pct": round(100 * confirmed / len(results), 1) if results else 0,
           "used_sources": used, "glossary": plain.GLOSSARY, "resolve": resolve}
    return run, metrics


def main(argv=None):
    a = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    a.add_argument("--client", required=True)
    a.add_argument("--lw")
    a.add_argument("--qbo")
    a.add_argument("--period", default=str(date.today()))
    a.add_argument("--mcp", action="store_true", help="JSON inputs are raw MCP tool output")
    args = a.parse_args(argv)
    try:
        lw_rows, qbo_rows = preflight(args.lw, args.qbo)
    except PreflightError as e:
        print(e)
        return 2
    try:
        d = client_dir(args.client)
    except PreflightError as e:
        print(e)
        return 2
    wb_path = os.path.join(d, "workbook.xlsx")
    try:
        run, metrics = analyse(lw_rows, qbo_rows, wb_path, args.client, args.period, args.mcp)
        written = workbook.update_workbook(wb_path, run, do_backup=False)  # apply_decisions already backed up
    except workbook.WorkbookDamaged as e:
        print(e)
        return 3
    summary.save_summary_csv(os.path.join(d, "summary.csv"), metrics)
    print(plain.headline({"decisions": workbook.open_count(written), "urgent": sum(1 for w in run["wrongs"] if w["severity"] in ("high", "urgent", "critical"))}))
    print(f"Workbook: {written}")
    if written != wb_path:
        print("Excel has the workbook open. Close it, copy the _pending file over the original, then run again.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
