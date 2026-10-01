#!/usr/bin/env python3
"""
Validate tax return review notes and build the review workbook.

The validation is the point. This script REFUSES to build a workbook when the
review notes fail the evidence discipline:

  - a row concluding 'agreed' or 'exception' with no source_document
  - a row concluding 'agreed' or 'exception' with no form/line reference
  - a 'must_fix_before_filing' item with no owner
  - a stated difference that does not equal per_return - expected
  - an unrecognised conclusion or severity value

Usage:
    python3 build_review.py --notes review_notes.csv \
        --client "Acme Holdings LLC" --return-type 1065 --tax-year 2025 \
        --checklist "Firm partnership checklist v3" \
        --out "Acme Holdings - 2025 1065 - Review Notes.xlsx"

Optional:
    --scope-received  "Return PDF, diagnostics, TB, FA schedule, PY return"
    --scope-missing   "Payment confirmations for Q3 estimate; K-1 from Meridian LP"
    --preparer "J. Alvarez"  --reviewer "S. Rao"
    --verdict "ready_subject_to_items"    ready | ready_subject_to_items | not_ready
    --force   build anyway, marked FAILED VALIDATION (for debugging only)

Notes CSV columns:
    ref, pass, area, form, line, item_tested, expected, per_return, difference,
    source_document, source_page, conclusion, severity, disposition, owner,
    due_date, notes

  conclusion : agreed | exception | open_item | n/a_explained
  severity   : must_fix_before_filing | should_fix | advisory | informational
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  install the openpyxl package")

MONEY = '#,##0.00;[Red](#,##0.00)'

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
ADV_FILL = PatternFill("solid", fgColor="DDEBF7")

CONCLUSIONS = {"agreed", "exception", "open_item", "n/a_explained"}
SEVERITIES = ["must_fix_before_filing", "should_fix", "advisory", "informational"]
SEV_FILL = {
    "must_fix_before_filing": BAD_FILL,
    "should_fix": WARN_FILL,
    "advisory": ADV_FILL,
    "informational": None,
}
VERDICTS = {
    "ready": "READY TO FILE",
    "ready_subject_to_items": "READY SUBJECT TO THE LISTED ITEMS",
    "not_ready": "NOT READY TO FILE",
}

COLS = [
    ("ref", "Ref", 8), ("pass", "Pass", 8), ("area", "Area", 24),
    ("form", "Form", 12), ("line", "Line", 10),
    ("item_tested", "Item tested", 44),
    ("expected", "Expected / per source", 18),
    ("per_return", "Per return", 16),
    ("difference", "Difference", 14),
    ("source_document", "Source document (evidence)", 34),
    ("source_page", "Pg", 6),
    ("conclusion", "Conclusion", 16),
    ("severity", "Severity", 22),
    ("disposition", "Disposition / what to do", 40),
    ("owner", "Owner", 14), ("due_date", "Due", 12),
    ("notes", "Notes", 46),
]


def dec(raw):
    s = str(raw or "").strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").strip()
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        return None          # non-numeric "expected" is legitimate (e.g. "Yes")
    return -v if neg else v


def load(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No review notes in {path}")
    out = []
    for i, r in enumerate(rows, start=1):
        rec = {k: " ".join(str(r.get(k, "") or "").split()) for k, _, _ in COLS}
        rec["_row"] = i
        rec["ref"] = rec["ref"] or f"R{i:03d}"
        rec["conclusion"] = rec["conclusion"].lower()
        rec["severity"] = rec["severity"].lower()
        out.append(rec)
    return out


def validate(rows: list[dict]) -> list[str]:
    errs: list[str] = []
    seen: dict[str, int] = {}

    for r in rows:
        n, ref = r["_row"], r["ref"]

        if ref in seen:
            errs.append(f"Row {n}: duplicate ref '{ref}' (also row {seen[ref]}).")
        seen[ref] = n

        if r["conclusion"] not in CONCLUSIONS:
            errs.append(
                f"Row {n} ({ref}): conclusion '{r['conclusion']}' is not one of "
                f"{sorted(CONCLUSIONS)}."
            )
        if r["severity"] and r["severity"] not in SEVERITIES:
            errs.append(
                f"Row {n} ({ref}): severity '{r['severity']}' is not one of "
                f"{SEVERITIES}."
            )
        if not r["item_tested"]:
            errs.append(f"Row {n} ({ref}): item_tested is empty. Say what was tested.")

        # --- the evidence discipline
        if r["conclusion"] in ("agreed", "exception"):
            if not r["source_document"]:
                errs.append(
                    f"Row {n} ({ref}): concluded '{r['conclusion']}' with no "
                    f"source_document. An item is not tested without evidence - cite "
                    f"the document, or change the conclusion to 'open_item'."
                )
            if not r["form"] and not r["line"]:
                errs.append(
                    f"Row {n} ({ref}): concluded '{r['conclusion']}' with no form or "
                    f"line reference. A reviewer must be able to find it on the return."
                )

        if r["severity"] == "must_fix_before_filing" and not r["owner"]:
            errs.append(
                f"Row {n} ({ref}): must_fix_before_filing with no owner. Every "
                f"blocking item needs a named person."
            )
        if r["conclusion"] == "exception" and not r["disposition"]:
            errs.append(
                f"Row {n} ({ref}): exception with no disposition. State what should "
                f"be done about it."
            )
        if r["conclusion"] == "n/a_explained" and not r["notes"]:
            errs.append(
                f"Row {n} ({ref}): concluded 'n/a_explained' but gave no explanation "
                f"in notes. 'Not applicable' is a conclusion and needs a reason."
            )

        # --- arithmetic
        exp, per, diff = dec(r["expected"]), dec(r["per_return"]), dec(r["difference"])
        if exp is not None and per is not None:
            computed = per - exp
            if diff is None:
                r["difference"] = f"{computed:.2f}"
            elif diff != computed:
                errs.append(
                    f"Row {n} ({ref}): stated difference {diff} does not equal "
                    f"per_return - expected ({per} - {exp} = {computed})."
                )
            if r["conclusion"] == "agreed" and computed != Decimal("0"):
                errs.append(
                    f"Row {n} ({ref}): concluded 'agreed' but per_return differs from "
                    f"expected by {computed}. That is an exception, not an agreement."
                )
    return errs


def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def sheet_summary(wb, rows, meta, verdict_key, scope_recv, scope_miss, valid):
    ws = wb.active
    ws.title = "Review Summary"
    for c, w in {1: 4, 2: 46, 3: 60, 4: 16, 5: 14, 6: 12}.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    r = 1

    def line(txt, *, bold=False, size=11, col=2, fill=None):
        nonlocal r
        c = ws.cell(row=r, column=col, value=txt)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        r += 1

    line("PREPARER REVIEW SUMMARY", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    if not valid:
        line("VALIDATION FAILED - this workbook does not meet the evidence standard. "
             "See the Validation tab.", bold=True, fill=BAD_FILL)
        r += 1

    sev = Counter(x["severity"] for x in rows)
    con = Counter(x["conclusion"] for x in rows)
    mustfix = [x for x in rows if x["severity"] == "must_fix_before_filing"]

    v = ws.cell(row=r, column=2, value="VERDICT: " + VERDICTS.get(verdict_key, "(not stated)"))
    v.font = OK_FONT if verdict_key == "ready" else BAD_FONT
    if verdict_key != "ready":
        v.fill = WARN_FILL if verdict_key == "ready_subject_to_items" else BAD_FILL
    r += 1
    line(f"{len(mustfix)} item(s) must be resolved before filing.", bold=True)
    r += 1

    line("SCOPE", bold=True, size=12)
    line(f"Items tested: {len(rows)}")
    line("Documents received: " + (scope_recv or "(not stated)"))
    mc = ws.cell(row=r, column=2,
                 value="Documents NOT received: " + (scope_miss or "(none noted)"))
    mc.font = Font(bold=bool(scope_miss))
    if scope_miss:
        mc.fill = WARN_FILL
        ws.cell(row=r, column=3,
                value="A review performed without the diagnostics list or prior-year "
                      "carryforwards is a limited-scope review and should say so."
                ).font = Font(italic=True, size=9)
    r += 2

    line("RESULTS", bold=True, size=12)
    for key in ("agreed", "exception", "open_item", "n/a_explained"):
        ws.cell(row=r, column=2, value=f"  {key}")
        ws.cell(row=r, column=4, value=con.get(key, 0))
        r += 1
    r += 1
    for s in SEVERITIES:
        ws.cell(row=r, column=2, value=f"  {s}")
        c = ws.cell(row=r, column=4, value=sev.get(s, 0))
        if s == "must_fix_before_filing" and sev.get(s):
            c.font = BAD_FONT
        r += 1
    r += 1

    if mustfix:
        line("MUST FIX BEFORE FILING", bold=True, size=12)
        for h, col in zip(["Ref", "Form / line", "Item", "Effect", "Owner", "Due"],
                          range(2, 8)):
            hc = ws.cell(row=r, column=col, value=h)
            hc.fill, hc.font = HDR_FILL, HDR_FONT
        r += 1
        for x in mustfix:
            ws.cell(row=r, column=2, value=x["ref"])
            ws.cell(row=r, column=3,
                    value=" ".join(p for p in (x["form"], x["line"]) if p))
            ws.cell(row=r, column=4, value=x["item_tested"])
            d = dec(x["difference"])
            dc = ws.cell(row=r, column=5, value=float(d) if d is not None else "")
            dc.number_format = MONEY
            ws.cell(row=r, column=6, value=x["owner"])
            ws.cell(row=r, column=7, value=x["due_date"])
            r += 1
        r += 1

    openitems = [x for x in rows if x["conclusion"] == "open_item"]
    if openitems:
        line("OPEN ITEMS - CLIENT INFORMATION NEEDED", bold=True, size=12)
        for x in openitems:
            line(f"  [{x['ref']}] {x['item_tested']} — {x['disposition'] or 'request from client'}")
        r += 1
        line("Next step: turn these into a plain-language client request.", bold=True)
        r += 1

    r += 1
    line("Prepared by / date:  __________________  ____________", bold=True)
    line("Reviewed by / date:  __________________  ____________", bold=True)


def sheet_notes(wb, rows):
    ws = wb.create_sheet("Review Notes")
    ws.append([label for _, label, _ in COLS])
    hdr(ws, len(COLS))
    for x in rows:
        out = []
        for key, _, _ in COLS:
            if key in ("expected", "per_return", "difference"):
                d = dec(x[key])
                out.append(float(d) if d is not None else (x[key] or None))
            else:
                out.append(x[key] or None)
        ws.append(out)

    idx = {k: i for i, (k, _, _) in enumerate(COLS)}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for k in ("expected", "per_return", "difference"):
            cell = row[idx[k]]
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY
        sv = row[idx["severity"]].value
        fill = SEV_FILL.get(sv) if sv else None
        if fill:
            row[idx["severity"]].fill = fill
        if row[idx["conclusion"]].value == "exception":
            row[idx["conclusion"]].font = BAD_FONT
        if not row[idx["source_document"]].value:
            row[idx["source_document"]].fill = WARN_FILL

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
    for i, (_, _, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_by_area(wb, rows):
    ws = wb.create_sheet("Coverage by Area")
    heads = ["Area", "Tested", "Agreed", "Exceptions", "Open items",
             "N/A explained", "Must fix"]
    ws.append(heads)
    hdr(ws, len(heads))
    agg = defaultdict(Counter)
    for x in rows:
        a = agg[x["area"] or "(unstated)"]
        a["n"] += 1
        a[x["conclusion"]] += 1
        if x["severity"] == "must_fix_before_filing":
            a["mustfix"] += 1
    for area, a in sorted(agg.items()):
        ws.append([area, a["n"], a["agreed"], a["exception"], a["open_item"],
                   a["n/a_explained"], a["mustfix"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[3].value:
            row[3].font = BAD_FONT
        if row[6].value:
            row[6].fill = BAD_FILL
    ws.freeze_panes = "A2"
    for c, w in {1: 34, 2: 10, 3: 10, 4: 12, 5: 12, 6: 15, 7: 11}.items():
        ws.column_dimensions[get_column_letter(c)].width = w


def sheet_validation(wb, errs):
    ws = wb.create_sheet("Validation")
    for c, w in {1: 4, 2: 118}.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.cell(row=1, column=2, value="EVIDENCE VALIDATION").font = Font(bold=True, size=14)
    if not errs:
        c = ws.cell(row=3, column=2,
                    value="PASSED - every tested item carries a form/line reference and "
                          "a source document, every blocking item has an owner, and all "
                          "stated differences foot.")
        c.font = OK_FONT
        return
    c = ws.cell(row=3, column=2, value=f"FAILED - {len(errs)} problem(s).")
    c.font, c.fill = BAD_FONT, BAD_FILL
    r = 5
    for e in errs:
        cell = ws.cell(row=r, column=2, value=e)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--notes", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--client", default="")
    ap.add_argument("--return-type", default="")
    ap.add_argument("--tax-year", default="")
    ap.add_argument("--checklist", default="")
    ap.add_argument("--scope-received", default="")
    ap.add_argument("--scope-missing", default="")
    ap.add_argument("--preparer", default="")
    ap.add_argument("--reviewer", default="")
    ap.add_argument("--verdict", default="", choices=[""] + list(VERDICTS))
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    rows = load(Path(args.notes))
    errs = validate(rows)
    valid = not errs

    mustfix = sum(1 for x in rows if x["severity"] == "must_fix_before_filing")
    openn = sum(1 for x in rows if x["conclusion"] == "open_item")
    verdict = args.verdict or (
        "not_ready" if mustfix else
        "ready_subject_to_items" if openn else "ready"
    )

    print("=" * 70)
    print("TAX RETURN REVIEW")
    print("=" * 70)
    print(f"Client      : {args.client or '(not stated)'}")
    print(f"Return      : {args.return_type or '(not stated)'}  TY {args.tax_year or '?'}")
    print(f"Checklist   : {args.checklist or '(not stated)'}")
    print(f"Items tested: {len(rows)}")
    con = Counter(x["conclusion"] for x in rows)
    print(f"  agreed {con['agreed']}   exception {con['exception']}   "
          f"open_item {con['open_item']}   n/a_explained {con['n/a_explained']}")
    print(f"Must fix before filing: {mustfix}")
    print(f"Verdict     : {VERDICTS[verdict]}")

    if errs:
        print(f"\nEVIDENCE VALIDATION FAILED - {len(errs)} problem(s):")
        for e in errs[:40]:
            print(f"  ! {e}")
        if len(errs) > 40:
            print(f"  ... and {len(errs) - 40} more")
        if not args.force:
            print("\n" + "=" * 70)
            print("WORKBOOK NOT WRITTEN.")
            print("An item concluded without a citation is not a tested item. Fix the")
            print("notes and re-run. Use --force only to inspect a marked-FAILED file.")
            print("=" * 70)
            return 1
    else:
        print("\nEvidence validation: PASSED")

    meta = {
        "Client": args.client or "(not stated)",
        "Return type": args.return_type or "(not stated)",
        "Tax year": args.tax_year or "(not stated)",
        "Checklist used": args.checklist or "(none stated - say which was used)",
        "Preparer": args.preparer or "(not stated)",
        "Reviewer": args.reviewer or "(not stated)",
        "Review date": datetime.now().strftime("%Y-%m-%d"),
    }

    wb = Workbook()
    sheet_summary(wb, rows, meta, verdict, args.scope_received, args.scope_missing, valid)
    sheet_notes(wb, rows)
    sheet_by_area(wb, rows)
    sheet_validation(wb, errs)
    wb.active = 0

    out = Path(args.out)
    if not valid:
        out = out.with_name(out.stem + " [FAILED VALIDATION]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\nWorkbook: {out}")
    return 0 if valid else 1


if __name__ == "__main__":
    sys.exit(main())
