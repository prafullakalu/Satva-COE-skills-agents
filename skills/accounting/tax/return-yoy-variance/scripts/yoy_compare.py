#!/usr/bin/env python3
"""
Year-over-year variance analysis with dual materiality and an explanation gate.

Flags a variance as requiring explanation when it breaches the dollar OR the
percent threshold (percent ignored below an absolute floor), plus five
always-flag conditions that are qualitatively material at any magnitude:

    sign flip · new line · disappeared line ·
    zero variance where the underlying activity changed · round-number plug

Reports explanation coverage and will not call the analysis complete until
every flagged item has a cause whose components sum to the variance.

Usage:
    python3 yoy_compare.py --current cy.csv --prior py.csv \
        --dollar-threshold 25000 --percent-threshold 10 --absolute-floor 5000 \
        --client "Acme Holdings LLC" --cy-label TY2025 --py-label TY2024 \
        --out "Acme - TY2025 vs TY2024 Variance.xlsx"

Input CSVs (current and prior, same shape):
    line_id     stable key used to match years              required
    caption     human-readable line name                    required
    amount      signed                                      required
    statement   e.g. Income Statement / Balance Sheet / Return   optional
    form        e.g. "1065 pg1"                             optional
    line_ref    e.g. "Ln 1c"                                optional
    group       subtotal grouping                           optional
    activity_changed   yes|no|blank  (current-year file only)    optional

--explanations CSV:
    line_id, cause, components, evidence, owner, notes
      components : semicolon-separated signed amounts that must sum to the
                   variance, e.g. "900000; 340000; -60000"
      evidence   : document reference, or the literal "client representation"
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required.  install the openpyxl package")

ZERO = Decimal("0.00")
MONEY = '#,##0.00;[Red](#,##0.00)'
PCT = '0.0"%"'

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
OK_FONT = Font(bold=True, color="006100")
BAD_FONT = Font(bold=True, color="9C0006")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
INFO_FILL = PatternFill("solid", fgColor="DDEBF7")

# Amounts that are suspiciously round on a line that should be computed.
# Roundness has to scale with magnitude or the test is pure noise: $5,360,000 being
# a multiple of 10,000 means nothing, but $100,000 being a multiple of 100,000 does.
# So the modulus must be a meaningful fraction (>=5%) of the amount itself.
ROUND_MODULI = (Decimal("1000000"), Decimal("500000"), Decimal("100000"),
                Decimal("50000"), Decimal("25000"), Decimal("10000"),
                Decimal("5000"), Decimal("1000"))
ROUND_MIN = Decimal("1000")
ROUND_SIGNIFICANCE = Decimal("20")   # modulus >= |amount| / 20


def dec(raw) -> Decimal | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if s in ("", "-", "--", "n/a", "N/A", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "").replace("$", "").replace("%", "").strip()
    if s.endswith("-"):
        neg, s = True, s[:-1]
    try:
        v = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Unparseable amount: {raw!r}")
    return -v if neg else v


def is_round(v: Decimal) -> str:
    a = abs(v)
    if a < ROUND_MIN:
        return ""
    need = a / ROUND_SIGNIFICANCE
    for m in ROUND_MODULI:
        if m >= need and a >= m and a % m == ZERO:
            return f"exact multiple of {m:,.0f}"
    return ""


def load(path: Path, label: str) -> dict[str, dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"No rows in {path}")
    out: dict[str, dict] = {}
    for i, r in enumerate(rows, start=1):
        lid = " ".join(str(r.get("line_id", "") or "").split())
        cap = " ".join(str(r.get("caption", "") or "").split())
        if not lid:
            lid = cap
        if not lid:
            raise ValueError(f"{label} row {i}: needs a line_id or a caption.")
        amt = dec(r.get("amount"))
        if amt is None:
            raise ValueError(f"{label} row {i} ({lid}): amount is required.")
        if lid in out:
            raise ValueError(
                f"{label}: duplicate line_id {lid!r}. Line IDs must be unique - "
                f"aggregate the duplicates or give them distinct IDs."
            )
        act = str(r.get("activity_changed", "") or "").strip().lower()
        out[lid] = {
            "line_id": lid, "caption": cap or lid, "amount": amt,
            "statement": " ".join(str(r.get("statement", "") or "").split()),
            "form": " ".join(str(r.get("form", "") or "").split()),
            "line_ref": " ".join(str(r.get("line_ref", "") or "").split()),
            "group": " ".join(str(r.get("group", "") or "").split()),
            "activity_changed": act in ("yes", "y", "true", "1"),
        }
    return out


def load_explanations(path: Path | None) -> dict[str, dict]:
    if not path or not path.exists():
        return {}
    out = {}
    with path.open(newline="", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            lid = " ".join(str(r.get("line_id", "") or "").split())
            if not lid:
                continue
            comps = []
            raw = str(r.get("components", "") or "")
            for part in re.split(r"[;|]", raw):
                part = part.strip()
                if part:
                    v = dec(part)
                    if v is not None:
                        comps.append(v)
            out[lid] = {
                "cause": " ".join(str(r.get("cause", "") or "").split()),
                "components_raw": raw.strip(),
                "components": comps,
                "evidence": " ".join(str(r.get("evidence", "") or "").split()),
                "owner": " ".join(str(r.get("owner", "") or "").split()),
                "notes": " ".join(str(r.get("notes", "") or "").split()),
            }
    return out


def compare(cy: dict, py: dict, dollar: Decimal, pct: Decimal,
            floor: Decimal) -> list[dict]:
    rows = []
    for lid in sorted(set(cy) | set(py)):
        c, p = cy.get(lid), py.get(lid)
        cy_amt = c["amount"] if c else None
        py_amt = p["amount"] if p else None
        src = c or p
        var = (cy_amt or ZERO) - (py_amt or ZERO)

        pct_var = None
        if py_amt not in (None, ZERO):
            pct_var = (var / abs(py_amt)) * Decimal(100)

        flags: list[str] = []
        if c and not p:
            flags.append("NEW line this year")
        if p and not c:
            flags.append("DISAPPEARED - present last year, absent this year")
        if cy_amt is not None and py_amt is not None:
            if (cy_amt > ZERO and py_amt < ZERO) or (cy_amt < ZERO and py_amt > ZERO):
                flags.append("SIGN FLIP - direction changed")
            if var == ZERO and c and c["activity_changed"]:
                flags.append("UNCHANGED but activity changed - stale rolled-forward figure?")
            if var == ZERO and cy_amt != ZERO and c and not c["activity_changed"]:
                flags.append("identical to prior year to the penny - confirm still correct")
        if cy_amt is not None:
            rnd = is_round(cy_amt)
            if rnd:
                flags.append(f"round number ({rnd}) - computed or plugged?")

        # dual materiality
        mat_d = abs(var) >= dollar
        mat_p = (pct_var is not None
                 and abs(pct_var) >= pct
                 and max(abs(cy_amt or ZERO), abs(py_amt or ZERO)) >= floor)
        material = mat_d or mat_p
        why = []
        if mat_d:
            why.append(f"|variance| >= {dollar:,.0f}")
        if mat_p:
            why.append(f"|%| >= {pct}% above {floor:,.0f} floor")

        needs = material or bool(flags)
        rows.append({
            "line_id": lid, "caption": src["caption"],
            "statement": src["statement"], "group": src["group"],
            "form": src["form"], "line_ref": src["line_ref"],
            "py": py_amt, "cy": cy_amt, "variance": var, "pct": pct_var,
            "material": material, "materiality_basis": "; ".join(why),
            "flags": flags, "needs_explanation": needs,
        })
    rows.sort(key=lambda r: abs(r["variance"]), reverse=True)
    return rows


def attach(rows: list[dict], expl: dict) -> tuple[list[dict], list[str]]:
    problems: list[str] = []
    for r in rows:
        e = expl.get(r["line_id"])
        r["cause"] = e["cause"] if e else ""
        r["components_raw"] = e["components_raw"] if e else ""
        r["evidence"] = e["evidence"] if e else ""
        r["owner"] = e["owner"] if e else ""
        r["notes"] = e["notes"] if e else ""
        r["explained"] = bool(e and e["cause"])
        r["components_tie"] = ""

        if not r["explained"]:
            continue
        if e["components"]:
            total = sum(e["components"], ZERO)
            diff = total - r["variance"]
            if diff == ZERO:
                r["components_tie"] = "ties"
            else:
                r["components_tie"] = f"OFF BY {diff:,.2f}"
                problems.append(
                    f"{r['line_id']} ({r['caption']}): components sum to {total:,.2f} "
                    f"but the variance is {r['variance']:,.2f} - {diff:,.2f} of the "
                    f"movement is unexplained. Do not absorb the remainder into 'other'."
                )
        elif r["material"]:
            r["components_tie"] = "no components given"
            problems.append(
                f"{r['line_id']} ({r['caption']}): material variance "
                f"{r['variance']:,.2f} has a cause but no quantified components. An "
                f"explanation that cannot be tested is a restatement of the variance."
            )
        if not r["evidence"]:
            problems.append(
                f"{r['line_id']} ({r['caption']}): explained with no evidence "
                f"reference. Cite a document, or state 'client representation' so the "
                f"weaker basis is visible to a reviewer."
            )
    for lid in sorted(set(expl) - {r["line_id"] for r in rows}):
        problems.append(
            f"Explanation supplied for line_id {lid!r}, which appears in neither year. "
            f"Likely a typo."
        )
    return rows, problems


# ------------------------------------------------------------------ workbook

def hdr(ws, n, row=1):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, w):
    for c, v in w.items():
        ws.column_dimensions[get_column_letter(c)].width = v


SCHED_COLS = [
    ("caption", "Caption", 42), ("statement", "Statement", 18),
    ("form", "Form", 12), ("line_ref", "Line", 10),
    ("py", "Prior year", 16), ("cy", "Current year", 16),
    ("variance", "Variance $", 16), ("pct", "Variance %", 12),
    ("material", "Material", 10), ("flagstr", "Flag reason", 46),
    ("cause", "Cause (causal + quantified)", 56),
    ("components_raw", "Components", 30),
    ("components_tie", "Tie", 14),
    ("evidence", "Evidence", 26), ("owner", "Owner", 14),
    ("line_id", "Line ID", 16),
]


def sheet_schedule(wb, rows, title="Variance Schedule", subset=None):
    ws = wb.create_sheet(title)
    ws.append([lbl for _, lbl, _ in SCHED_COLS])
    hdr(ws, len(SCHED_COLS))
    data = subset if subset is not None else rows
    for r in data:
        r["flagstr"] = "; ".join(r["flags"])
        out = []
        for key, _, _ in SCHED_COLS:
            v = r.get(key)
            if key in ("py", "cy", "variance"):
                out.append(float(v) if v is not None else None)
            elif key == "pct":
                out.append(float(v) if v is not None else None)
            elif key == "material":
                out.append("YES" if v else "")
            else:
                out.append(v or None)
        ws.append(out)

    idx = {k: i for i, (k, _, _) in enumerate(SCHED_COLS)}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for k in ("py", "cy", "variance"):
            row[idx[k]].number_format = MONEY
        row[idx["pct"]].number_format = PCT
        if row[idx["material"]].value == "YES":
            row[idx["material"]].font = BAD_FONT
        if row[idx["flagstr"]].value:
            row[idx["flagstr"]].fill = WARN_FILL
        tie = row[idx["components_tie"]].value
        if tie and tie != "ties":
            row[idx["components_tie"]].fill = BAD_FILL
            row[idx["components_tie"]].font = BAD_FONT
        elif tie == "ties":
            row[idx["components_tie"]].font = OK_FONT
        if (row[idx["material"]].value == "YES" or row[idx["flagstr"]].value) \
                and not row[idx["cause"]].value:
            row[idx["cause"]].fill = BAD_FILL
            row[idx["cause"]].value = "*** UNEXPLAINED ***"
        if str(row[idx["evidence"]].value or "").lower() == "client representation":
            row[idx["evidence"]].fill = INFO_FILL

    if ws.max_row == 1:
        ws.cell(row=2, column=1, value="None.")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SCHED_COLS))}{max(ws.max_row, 2)}"
    widths(ws, {i: w for i, (_, _, w) in enumerate(SCHED_COLS, start=1)})


def sheet_summary(wb, rows, meta, stats, problems):
    ws = wb.active
    ws.title = "Summary"
    widths(ws, {1: 4, 2: 50, 3: 20, 4: 20, 5: 62})
    r = 1

    def line(txt, *, bold=False, size=11, fill=None, col=2):
        nonlocal r
        c = ws.cell(row=r, column=col, value=txt)
        c.font = Font(bold=bold, size=size)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
        r += 1

    line("YEAR-OVER-YEAR VARIANCE ANALYSIS", bold=True, size=14)
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=2, value=k).font = Font(bold=True)
        ws.cell(row=r, column=3, value=v)
        r += 1
    r += 1

    complete = stats["needs"] > 0 and stats["explained"] == stats["needs"] and not problems
    if stats["needs"] == 0:
        line("No variance breached materiality and no always-flag condition arose.",
             bold=True, fill=INFO_FILL)
    elif complete:
        line("ANALYSIS COMPLETE - every flagged variance has a quantified, "
             "evidenced cause.", bold=True, fill=None)
        ws.cell(row=r - 1, column=2).font = OK_FONT
    else:
        line(f"ANALYSIS INCOMPLETE - {stats['needs'] - stats['explained']} flagged "
             f"item(s) unexplained"
             + (f", {len(problems)} explanation problem(s)" if problems else "")
             + ". See Unexplained tab.", bold=True, fill=BAD_FILL)
    r += 1

    line("COVERAGE", bold=True, size=12)
    for label, val in [
        ("Lines compared", stats["total"]),
        ("Requiring explanation", stats["needs"]),
        ("Explained", stats["explained"]),
        ("Explanation rate (count)", f"{stats['rate_count']:.0f}%"),
        ("Flagged variance $ (absolute)", None),
        ("Explained variance $ (absolute)", None),
        ("Explanation rate (dollars)", f"{stats['rate_dollars']:.0f}%"),
    ]:
        ws.cell(row=r, column=2, value="  " + label)
        if label == "Flagged variance $ (absolute)":
            c = ws.cell(row=r, column=3, value=float(stats["flagged_dollars"]))
            c.number_format = MONEY
        elif label == "Explained variance $ (absolute)":
            c = ws.cell(row=r, column=3, value=float(stats["explained_dollars"]))
            c.number_format = MONEY
        else:
            ws.cell(row=r, column=3, value=val)
        r += 1
    r += 1

    top = [x for x in rows if x["variance"] != ZERO][:10]
    if top:
        line("TOP 10 VARIANCES BY ABSOLUTE DOLLAR", bold=True, size=12)
        for h, col in zip(["Caption", "Prior", "Current", "Variance", "Cause"],
                          (2, 3, 4, 5, 6)):
            hc = ws.cell(row=r, column=col, value=h)
            hc.fill, hc.font = HDR_FILL, HDR_FONT
        r += 1
        for x in top:
            ws.cell(row=r, column=2, value=x["caption"])
            for col, key in ((3, "py"), (4, "cy"), (5, "variance")):
                v = x[key]
                c = ws.cell(row=r, column=col, value=float(v) if v is not None else None)
                c.number_format = MONEY
            cc = ws.cell(row=r, column=6, value=x["cause"] or "*** UNEXPLAINED ***")
            if not x["cause"]:
                cc.fill, cc.font = BAD_FILL, BAD_FONT
            r += 1
        r += 1

    line("RELATIONSHIP TESTS TO PERFORM (see Relationship Tests tab)", bold=True, size=12)
    for t in [
        "Gross margin % by year",
        "Revenue vs. payroll and headcount",
        "Revenue vs. receivables and DSO",
        "Fixed-asset additions vs. depreciation expense (stale-schedule detector)",
        "Debt balances vs. interest expense (implied rate)",
        "Officer compensation vs. distributions (S corps)",
        "Book income vs. taxable income, by component",
        "Effective tax rate, reconciled (top-down check)",
    ]:
        line("  - " + t)
    r += 1
    line("An explanation is causal and quantified. 'Timing' and 'higher volume' are "
         "not explanations - they restate the variance.", bold=True)


def sheet_flags(wb, rows):
    subset = [r for r in rows if r["flags"]]
    sheet_schedule(wb, rows, "Always-Flag Exceptions", subset)


def sheet_unexplained(wb, rows, problems):
    subset = [r for r in rows if r["needs_explanation"] and not r["explained"]]
    sheet_schedule(wb, rows, "Unexplained", subset)
    ws = wb["Unexplained"]
    if problems:
        r = ws.max_row + 3
        c = ws.cell(row=r, column=1, value="EXPLANATION PROBLEMS")
        c.font = Font(bold=True, size=12)
        for p in problems:
            r += 1
            cell = ws.cell(row=r, column=1, value=p)
            cell.font = BAD_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_relationships(wb, cy_label, py_label):
    ws = wb.create_sheet("Relationship Tests")
    heads = ["Test", py_label, cy_label, "Movement", "Explanation", "Evidence"]
    ws.append(heads)
    hdr(ws, len(heads))
    for t in [
        "Gross margin %",
        "Payroll as % of revenue",
        "Revenue per employee",
        "DSO (days sales outstanding)",
        "Inventory turns",
        "Fixed-asset additions",
        "Depreciation expense",
        "Depreciation as % of gross fixed assets",
        "Debt balance (period end)",
        "Interest expense",
        "Implied interest rate %",
        "Officer compensation",
        "Distributions to owners",
        "Book income",
        "Taxable income",
        "Total book-to-tax difference",
        "Effective tax rate %",
    ]:
        ws.append([t, None, None, None, None, None])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (1, 2, 3):
            row[i].number_format = MONEY
    ws.freeze_panes = "B2"
    widths(ws, {1: 42, 2: 18, 3: 18, 4: 16, 5: 62, 6: 26})


# ------------------------------------------------------------------ main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--current", required=True)
    ap.add_argument("--prior", required=True)
    ap.add_argument("--explanations")
    ap.add_argument("--dollar-threshold", default="25000")
    ap.add_argument("--percent-threshold", default="10")
    ap.add_argument("--absolute-floor", default="5000",
                    help="below this magnitude, percent variance is ignored as noise")
    ap.add_argument("--client", default="")
    ap.add_argument("--cy-label", default="Current year")
    ap.add_argument("--py-label", default="Prior year")
    ap.add_argument("--basis-note", default="",
                    help="normalization applied, e.g. 'PY restated for amended return'")
    ap.add_argument("--out", required=True)
    ap.add_argument("--unexplained-out", default="unexplained.csv")
    args = ap.parse_args()

    dollar = dec(args.dollar_threshold) or ZERO
    pct = dec(args.percent_threshold) or ZERO
    floor = dec(args.absolute_floor) or ZERO

    cy = load(Path(args.current), "current")
    py = load(Path(args.prior), "prior")
    rows = compare(cy, py, dollar, pct, floor)
    expl = load_explanations(Path(args.explanations) if args.explanations else None)
    rows, problems = attach(rows, expl)

    needs = [r for r in rows if r["needs_explanation"]]
    explained = [r for r in needs if r["explained"]]
    flagged_dollars = sum((abs(r["variance"]) for r in needs), ZERO)
    explained_dollars = sum((abs(r["variance"]) for r in explained), ZERO)
    stats = {
        "total": len(rows), "needs": len(needs), "explained": len(explained),
        "rate_count": (len(explained) / len(needs) * 100) if needs else 100.0,
        "flagged_dollars": flagged_dollars,
        "explained_dollars": explained_dollars,
        "rate_dollars": (float(explained_dollars / flagged_dollars) * 100)
                        if flagged_dollars else 100.0,
    }

    # worklist for the preparer
    unexp = [r for r in needs if not r["explained"]]
    if unexp:
        with Path(args.unexplained_out).open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["line_id", "caption", "prior", "current", "variance", "pct",
                        "flag_reason", "cause", "components", "evidence", "owner", "notes"])
            for r in unexp:
                w.writerow([
                    r["line_id"], r["caption"],
                    f"{r['py']:.2f}" if r["py"] is not None else "",
                    f"{r['cy']:.2f}" if r["cy"] is not None else "",
                    f"{r['variance']:.2f}",
                    f"{r['pct']:.1f}" if r["pct"] is not None else "",
                    "; ".join(r["flags"]) or r["materiality_basis"],
                    "", "", "", "", "",
                ])

    meta = {
        "Client": args.client or "(not stated)",
        "Current year": args.cy_label,
        "Prior year": args.py_label,
        "Dollar threshold": float(dollar),
        "Percent threshold": f"{pct}%",
        "Absolute floor (percent test ignored below)": float(floor),
        "Comparability / normalization": args.basis_note or "(none stated - confirm same "
                                                            "entity, method, period length, "
                                                            "and chart of accounts)",
    }

    wb = Workbook()
    sheet_summary(wb, rows, meta, stats, problems)
    sheet_schedule(wb, rows)
    sheet_flags(wb, rows)
    sheet_relationships(wb, args.cy_label, args.py_label)
    sheet_unexplained(wb, rows, problems)
    wb.active = 0

    complete = (stats["needs"] == 0) or (stats["explained"] == stats["needs"] and not problems)
    out = Path(args.out)
    if not complete:
        out = out.with_name(out.stem + " [INCOMPLETE]" + out.suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    # ---- console
    print("=" * 72)
    print(f"YEAR-OVER-YEAR VARIANCE:  {args.py_label} -> {args.cy_label}")
    print("=" * 72)
    print(f"Materiality: |var| >= {dollar:,.0f}  OR  |%| >= {pct}%  "
          f"(percent ignored below {floor:,.0f})")
    print(f"Lines compared        : {stats['total']}")
    print(f"Requiring explanation : {stats['needs']}")
    print(f"Explained             : {stats['explained']}  "
          f"({stats['rate_count']:.0f}% of items, {stats['rate_dollars']:.0f}% of dollars)")

    flags = [r for r in rows if r["flags"]]
    if flags:
        print(f"\nAlways-flag exceptions ({len(flags)}):")
        for r in flags[:20]:
            print(f"  ! {r['caption'][:44]:<44} {r['variance']:>14,.2f}  "
                  f"{'; '.join(r['flags'])[:60]}")
        if len(flags) > 20:
            print(f"  ... and {len(flags) - 20} more")

    if unexp:
        print(f"\nUNEXPLAINED ({len(unexp)}):")
        for r in unexp[:20]:
            print(f"  ? {r['caption'][:44]:<44} {r['variance']:>14,.2f}")
        if len(unexp) > 20:
            print(f"  ... and {len(unexp) - 20} more")
        print(f"\nWorklist written to {args.unexplained_out}. Fill in cause, "
              f"components, evidence, owner and re-run with --explanations.")

    if problems:
        print(f"\nEXPLANATION PROBLEMS ({len(problems)}):")
        for p in problems[:20]:
            print(f"  ! {p}")

    print(f"\nStatus  : {'COMPLETE' if complete else 'INCOMPLETE'}")
    print(f"Workbook: {out}")
    return 0 if complete else 1


if __name__ == "__main__":
    sys.exit(main())
